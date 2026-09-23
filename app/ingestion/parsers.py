"""
Document parsers. Each takes a file path and returns a list of
(text, page_or_slide_number) tuples so metadata can track where
a chunk came from.
"""
from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation


def parse_pdf(path: str) -> list[tuple[str, int]]:
    reader = PdfReader(path)
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append((text, i))
    return pages


def parse_pptx(path: str) -> list[tuple[str, int]]:
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        parts.append(line)
            # speaker notes are often where the real explanation lives
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                parts.append(f"[Speaker notes]: {notes}")
        text = "\n".join(parts)
        if text.strip():
            slides.append((text, i))
    return slides


def parse_markdown(path: str) -> list[tuple[str, int]]:
    text = Path(path).read_text(encoding="utf-8")
    # markdown has no page concept, treat the whole file as unit 1;
    # the chunker will split it further downstream.
    return [(text, 1)] if text.strip() else []


def parse_document(path: str) -> list[tuple[str, int]]:
    """Dispatches to the right parser based on file extension."""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    elif ext == ".pptx":
        return parse_pptx(path)
    elif ext in (".md", ".txt"):
        return parse_markdown(path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
